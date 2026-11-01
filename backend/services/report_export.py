"""One selection/query shared by both formats; lossless data travels with each file."""
from datetime import date, datetime, time
from functools import lru_cache
from io import BytesIO
from pathlib import Path
import tempfile
import json

from fastapi import HTTPException
from sqlalchemy import select, func

from services import report_service as reports

FONT = Path(__file__).resolve().parents[2] / "frontend/assets/fonts/vazirmatn/Vazirmatn-Variable.woff2"
MAX_ROWS = 50000


def display(value):
    if value is None: return "NULL"
    if value == "": return '""'
    if isinstance(value, (dict, list, bool)): return reports.canonical(value)
    return str(reports.scalar(value))


@lru_cache(maxsize=1)
def font_bytes():
    from fontTools.ttLib import TTFont
    from fontTools.varLib.instancer import instantiateVariableFont
    font = TTFont(FONT)
    if "fvar" in font:
        font = instantiateVariableFont(font, {a.axisTag: a.defaultValue for a in font["fvar"].axes})
    font.flavor = None
    out = BytesIO()
    font.save(out)
    return out.getvalue()


def export(db, name, request, fmt, destination):
    ds = reports.dataset(db, name)
    names = reports.selected(ds, request)
    reports.snapshot(db)
    total = db.scalar(select(func.count()).select_from(ds.source).where(*reports.conditions(ds, request)))
    if total > MAX_ROWS:
        raise HTTPException(422, "حداکثر ۵۰٬۰۰۰ رکورد در هر خروجی؛ بازه یا فیلتر را محدود کنید.")
    info = reports.schema(db, name)
    descriptors = {c["name"]: c for c in info["columns"]}
    columns = [descriptors[n] for n in names]
    statement = reports.query(ds, request, names).execution_options(yield_per=200)
    # Finish the read quickly before PDF shaping/ZIP compression. Rendering must
    # not hold SQLite read locks and block normal edits for several seconds.
    with tempfile.SpooledTemporaryFile(max_size=1024*1024, mode="w+", encoding="utf-8") as spool:
        result = db.execute(statement).mappings()
        try:
            for row in result:
                spool.write(reports.canonical(dict(row)) + "\n")
        finally:
            result.close()
            db.rollback()  # Read-only snapshot; release its connection and locks.
        spool.seek(0)
        rows = (json.loads(line) for line in spool)
        if fmt == "xlsx":
            write_xlsx(rows, columns, info["title"], destination)
        else:
            write_pdf(rows, columns, info["title"], destination)
    return total


def write_xlsx(rows, columns, title, destination):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    wb = Workbook(write_only=True)
    sheet = wb.create_sheet("گزارش")
    raw = wb.create_sheet("داده خام")
    raw.sheet_state = "hidden"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = "8"  # OOXML A3
    sheet.page_setup.fitToWidth = 1
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 72 if any(c["related"] for c in columns) else 48
    for i, col in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 38 if col["type"] in {"json", "text"} else 25
    font = Font(name="Vazirmatn", size=11)
    alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)

    def cell(ws, value, header=False):
        c = WriteOnlyCell(ws, value=value)
        # Never interpret database strings (including =, +, @) as formulas.
        if isinstance(value, str): c.data_type = "s"
        c.font = Font(name="Vazirmatn", size=11, bold=True, color="FFFFFF") if header else font
        c.alignment = alignment
        if header: c.fill = PatternFill("solid", fgColor="1E40AF")
        return c

    def raw_entry(kind, index, value):
        text = reports.canonical(value)
        for part, start in enumerate(range(0, len(text), 30000)):
            raw.append([cell(raw, v) for v in [kind, index, part, text[start:start+30000]]])

    raw_entry("schema", 0, dict(title=title, columns=columns, encoding="JSON", version=1))
    sheet.append([cell(sheet, c["label"] + "\n" + c["name"], True) for c in columns])
    count = 0
    for count, row in enumerate(rows, 1):
        values = dict(row)
        raw_entry("record", count, values)
        result = []
        for col in columns:
            v = row[col["name"]]
            if v is not None and col["type"] in {"date", "datetime", "time"}:
                v = {"date": date, "datetime": datetime, "time": time}[col["type"]].fromisoformat(v)
            if isinstance(v, (dict, list)): v = reports.canonical(v)
            if isinstance(v, int) and not isinstance(v, bool) and abs(v) >= 10**15: v = str(v)
            if isinstance(v, datetime) and v.tzinfo is not None: v = v.replace(tzinfo=None)
            note = None
            if isinstance(v, str):
                clean = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m[0]):04x}", v)
                if len(clean) > 32767:
                    clean = clean[:32600] + "\n[متن کامل در برگه داده خام]"
                    note = "متن از ظرفیت سلول Excel بیشتر است. مقدار کامل در برگه مخفی داده خام، رکورد " + str(count)
                v = clean
            c = cell(sheet, v)
            if isinstance(v, datetime): c.number_format = "yyyy-mm-dd hh:mm:ss.000000"
            elif isinstance(v, date): c.number_format = "yyyy-mm-dd"
            elif isinstance(v, time): c.number_format = "hh:mm:ss.000000"
            elif isinstance(v, int) and not isinstance(v, bool): c.number_format = "#,##0"
            elif isinstance(v, float): c.number_format = "#,##0.########"
            if note: c.comment = Comment(note, "Karbon")
            result.append(c)
        sheet.append(result)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, count+1)}"
    wb.properties.title = title
    wb.properties.description = "مقادیر کامل، NULL، رشته خالی، JSON و زمان اصلی در برگه مخفی «داده خام» حفظ شده‌اند."
    wb.save(destination)


def write_pdf(rows, columns, title, destination):
    from fpdf import FPDF
    from fpdf.enums import MethodReturnValue

    class ReportPDF(FPDF):
        def footer(self):
            self.set_y(-9)
            self.set_font("Vazirmatn", size=8)
            self.set_text_color(90, 100, 115)
            self.cell(0, 4, text=f"Karbon | {self.page_no()}", align="C")

    with tempfile.TemporaryDirectory(prefix="karbon-report-font-") as directory:
        font_path = Path(directory) / "Vazirmatn.ttf"
        font_path.write_bytes(font_bytes())
        pdf = ReportPDF(orientation="L", format="A3")
        pdf.set_margins(12, 12, 12)
        pdf.set_auto_page_break(False)
        pdf.add_font("Vazirmatn", fname=str(font_path))
        pdf.set_font("Vazirmatn", size=9)
        pdf.set_text_shaping(True)
        pdf.set_title(title)
        pdf.set_lang("fa-IR")
        pdf.set_creator("Karbon")
        pdf.set_display_mode("fullwidth")
        pdf.set_compression(True)
        # Wide reports use horizontal bands, never shrink text to fit every field.
        bands = [columns[i:i+5] for i in range(0, len(columns), 5)]
        with (Path(directory) / "report-data.json").open("w+", encoding="utf-8") as raw:
            raw.write('{"columns":' + reports.canonical(columns) + ',"records":[')
            first = True
            index = 0
            while True:
                batch = []
                for _ in range(100):
                    row = next(rows, None)
                    if row is None: break
                    values = dict(row)
                    raw.write(("" if first else ",") + reports.canonical(values))
                    first = False
                    batch.append(values)
                if not batch and index: break
                for band_index, band in enumerate(bands, 1):
                    width = pdf.epw / len(band)
                    def header():
                        pdf.add_page()
                        pdf.set_text_color(30, 64, 175)
                        pdf.set_font("Vazirmatn", size=13)
                        pdf.cell(0, 8, text=title, align="R", new_x="LMARGIN", new_y="NEXT")
                        pdf.set_font("Vazirmatn", size=9)
                        pdf.set_text_color(65, 75, 90)
                        pdf.cell(0, 6, text=f"بخش ستون‌ها: {band_index}/{len(bands)} | رکوردهای {index+1} تا {index+len(batch)}", align="R", new_x="LMARGIN", new_y="NEXT")
                        y = pdf.y
                        pdf.set_fill_color(30, 64, 175)
                        pdf.set_text_color(255)
                        for i, c in enumerate(band):
                            pdf.set_xy(pdf.w-pdf.r_margin-(i+1)*width, y)
                            pdf.rect(pdf.x, y, width, 15, style="F")
                            pdf.multi_cell(width, 5, text=c["label"]+"\n"+c["name"], align="R", wrapmode="CHAR")
                        pdf.set_xy(pdf.l_margin, y+15)
                        pdf.set_text_color(15, 23, 42)
                    header()
                    if not batch:
                        pdf.cell(0, 8, text="رکوردی وجود ندارد.", align="R")
                    for row_number, row in enumerate(batch):
                        # Lay out every value fully; even a single very long JSON/text
                        # cell can continue onto another page with repeated headers.
                        lines = []
                        for c in band:
                            pdf.set_x(pdf.l_margin)
                            lines.append(pdf.multi_cell(width-4, 5, text=display(row[c["name"]]),
                                dry_run=True, output=MethodReturnValue.LINES, wrapmode="CHAR"))
                        remaining = max(map(len, lines))
                        offset = 0
                        while offset < remaining:
                            available = int((pdf.h-15-pdf.y-4)/5)
                            if available < 1:
                                header()
                                continue
                            take = min(available, remaining-offset)
                            height = take*5+4
                            y = pdf.y
                            pdf.set_fill_color(*( (241,245,249) if row_number%2 == 0 else (255,255,255)))
                            pdf.set_draw_color(210, 218, 230)
                            for i, content in enumerate(lines):
                                x = pdf.w-pdf.r_margin-(i+1)*width
                                pdf.rect(x, y, width, height, style="DF")
                                pdf.set_xy(x+2, y+2)
                                text = "\n".join(content[offset:offset+take])
                                if text: pdf.multi_cell(width-4, 5, text=text, align="R", wrapmode="CHAR")
                            pdf.set_xy(pdf.l_margin, y+height)
                            offset += take
                index += len(batch)
                if not batch: break
            raw.write("]}")
            raw.flush()
            pdf.embed_file(file_path=raw.name, desc="Lossless selected fields and records (JSON)")
            pdf.output(destination)
