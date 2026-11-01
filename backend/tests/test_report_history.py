"""Durable report history on a disposable database, including immutable bytes and ownership."""
from contextlib import closing
from io import BytesIO
import json
import os
import sqlite3
import subprocess
import sys
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from pypdf import PdfReader
import test_customers as customers


class ReportHistoryTests(unittest.TestCase):
    setUpClass = classmethod(customers.CustomersIntegrationTests.setUpClass.__func__)
    tearDownClass = classmethod(customers.CustomersIntegrationTests.tearDownClass.__func__)
    request = classmethod(customers.CustomersIntegrationTests.request.__func__)

    def setUp(self):
        with closing(sqlite3.connect(self.database)) as db, db:
            db.execute("DELETE FROM generated_reports")
            db.execute("DELETE FROM customers")
        status, self.customer = self.request("POST", "/customers", {"first_name": "مشتری اصلی", "last_name": "آزمایش", "phone_numbers": []})
        self.assertEqual(status, 201)

    def binary(self, path, body=None, token=None):
        request = Request(self.base + path, data=None if body is None else json.dumps(body).encode(),
                          headers={"Authorization": "Bearer " + (token or self.token), "Content-Type": "application/json"})
        try:
            response = urlopen(request, timeout=30)
        except HTTPError as error:
            response = error
        with response:
            return response.status, response.read(), response.headers

    def test_saved_formats_metadata_and_download_survive_business_edits(self):
        self.assertEqual(self.request("GET", "/reports/history")[1], {"total": 0, "items": []})
        original = {}
        for fmt in ["xlsx", "pdf"]:
            name = "گزارش مشتریان <آزمایش> " + fmt
            status, content, headers = self.binary(f"/reports/customers/export/{fmt}", {
                "columns": ["first_name", "last_name"], "ids": [self.customer["id"]], "report_name": name})
            self.assertEqual(status, 200, content)
            self.assertEqual(headers["Cache-Control"], "no-store")
            status, history = self.request("GET", "/reports/history?limit=1")
            report = history["items"][0]
            original[report["id"]] = content
            self.assertEqual(history["total"], len(original))
            self.assertEqual(report["name"], name)
            self.assertEqual(report["format"], fmt)
            self.assertEqual(report["record_count"], 1)
            self.assertEqual(report["column_count"], 2)
            self.assertEqual(report["size_bytes"], len(content))
            self.assertTrue(report["created_at"].endswith("+00:00"))
            self.assertNotIn("content", report)
            self.assertNotIn("parameters", report)
        older = self.request("GET", "/reports/history?offset=1&limit=1")[1]["items"][0]
        self.assertEqual(older["format"], "xlsx")
        self.assertEqual(self.request("GET", "/reports/history?offset=2")[1]["items"], [])
        with closing(sqlite3.connect(self.database)) as db, db:
            db.execute("UPDATE customers SET first_name='changed' WHERE id=?", (self.customer["id"],))
        for report_id, expected in original.items():
            status, actual, headers = self.binary(f"/reports/history/{report_id}/download")
            self.assertEqual(status, 200)
            self.assertEqual(actual, expected)
            self.assertEqual(headers["Cache-Control"], "no-store")
            self.assertIn("attachment", headers["Content-Disposition"])
        # Re-run application startup in a fresh process against this existing DB.
        # Its create_all/migrations must preserve every saved export.
        subprocess.run([sys.executable, "-c", "import app"], cwd=customers.ROOT / "backend",
                       env=dict(os.environ, KARBON_DATABASE_PATH=self.database,
                                KARBON_JWT_SECRET="isolated-test-secret-only"), check=True, timeout=30)
        with closing(sqlite3.connect(self.database)) as db:
            self.assertEqual(dict(db.execute("SELECT id, content FROM generated_reports")), original)
        pdf = PdfReader(BytesIO(original[max(original)]))
        self.assertEqual(json.loads(pdf.attachments["report-data.json"][0])["records"][0]["first_name"], "مشتری اصلی")
        workbook = load_workbook(BytesIO(original[min(original)]))
        self.assertEqual(workbook["گزارش"].cell(2, 1).value, "مشتری اصلی")
        workbook.close()

    def test_authentication_ownership_and_invalid_exports(self):
        status, _, _ = self.binary("/reports/customers/export/xlsx", {"columns": ["id"]})
        self.assertEqual(status, 200)
        report_id = self.request("GET", "/reports/history")[1]["items"][0]["id"]
        self.assertEqual(self.request("GET", f"/reports/history/{report_id}", authenticated=False)[0], 401)
        self.assertEqual(self.request("DELETE", f"/reports/history/{report_id}", authenticated=False)[0], 401)
        self.assertEqual(self.request("GET", "/reports/history", authenticated=False)[0], 401)
        self.assertEqual(self.request("GET", f"/reports/history/{report_id}/download", authenticated=False)[0], 401)
        self.assertEqual(self.request("GET", "/reports/history/999999/download")[0], 404)
        with closing(sqlite3.connect(self.database)) as db, db:
            db.execute("INSERT OR IGNORE INTO users(username,password_hash,auth_version) SELECT 'report-other',password_hash,0 FROM users WHERE username='customers-test'")
        other = self.request("POST", "/login", {"username": "report-other", "password": "test-password"}, authenticated=False)[1]["token"]
        status, raw, _ = self.binary("/reports/history", token=other)
        self.assertEqual(json.loads(raw), {"total": 0, "items": []})
        self.assertEqual(self.binary(f"/reports/history/{report_id}/download", token=other)[0], 404)
        for method in ["GET", "DELETE"]:
            req = Request(self.base + f"/reports/history/{report_id}", method=method,
                          headers={"Authorization": "Bearer " + other})
            with self.assertRaises(HTTPError) as error:
                urlopen(req, timeout=10)
            self.assertEqual(error.exception.code, 404)
            error.exception.close()
        self.assertEqual(self.request("GET", "/reports/generated_reports/schema")[0], 404)
        before = self.request("GET", "/reports/history")[1]["total"]
        self.assertEqual(self.binary("/reports/customers/export/pdf", {"columns": ["missing"]})[0], 422)
        self.assertEqual(self.request("GET", "/reports/history")[1]["total"], before)
        self.assertEqual(self.request("GET", "/reports/history?limit=201")[0], 422)

    def test_details_and_delete_only_the_requested_report(self):
        for fmt in ["pdf", "xlsx"]:
            self.assertEqual(self.binary(f"/reports/customers/export/{fmt}", {"columns": ["id"]})[0], 200)
        history = self.request("GET", "/reports/history")[1]
        report, remaining = history["items"]
        status, detail = self.request("GET", f'/reports/history/{report["id"]}')
        self.assertEqual(status, 200)
        self.assertEqual(detail, report)
        self.assertNotIn("content", detail)
        self.assertEqual(self.request("DELETE", f'/reports/history/{report["id"]}')[0], 204)
        for suffix in ["", "/download"]:
            self.assertEqual(self.request("GET", f'/reports/history/{report["id"]}{suffix}')[0], 404)
        self.assertEqual(self.request("DELETE", f'/reports/history/{report["id"]}')[0], 404)
        self.assertEqual(self.request("GET", "/reports/history")[1], {"total": 1, "items": [remaining]})
        with closing(sqlite3.connect(self.database)) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM generated_reports").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT count(*) FROM customers").fetchone()[0], 1)

    def test_save_failure_leaves_no_partial_history_and_can_retry(self):
        with closing(sqlite3.connect(self.database)) as db, db:
            db.execute("CREATE TRIGGER fail_report_history BEFORE INSERT ON generated_reports BEGIN SELECT RAISE(ABORT, 'test save failure'); END")
        try:
            self.assertEqual(self.binary("/reports/customers/export/xlsx", {"columns": ["id"]})[0], 500)
            self.assertEqual(self.request("GET", "/reports/history")[1]["total"], 0)
        finally:
            with closing(sqlite3.connect(self.database)) as db, db:
                db.execute("DROP TRIGGER fail_report_history")
        self.assertEqual(self.binary("/reports/customers/export/xlsx", {"columns": ["id"]})[0], 200)
        self.assertEqual(self.request("GET", "/reports/history")[1]["total"], 1)



if __name__ == "__main__":
    unittest.main()
