"""HTTP, SQLite, XLSX and PDF integration on disposable data only.

Requires report dependencies plus pypdf for attachment/structure verification.
Run: python3 -m unittest discover -s backend/tests -p test_reports.py -v
"""
from collections import defaultdict
from contextlib import closing
from io import BytesIO
import json
import concurrent.futures
from pathlib import Path
import sqlite3
import time
import unittest
from urllib.request import Request, urlopen
from urllib.parse import quote
from urllib.error import HTTPError

from openpyxl import load_workbook
from pypdf import PdfReader
import test_customers as customers


TABLES = ["customers", "bank_directory", "bank_accounts", "expense_categories", "products",
          "documents", "document_accounts", "financial_transactions", "orders", "order_items",
          "order_payments", "inventory_options", "members", "attendance", "tasks"]


class ReportsIntegrationTests(unittest.TestCase):
    setUpClass = classmethod(customers.CustomersIntegrationTests.setUpClass.__func__)
    tearDownClass = classmethod(customers.CustomersIntegrationTests.tearDownClass.__func__)
    request = classmethod(customers.CustomersIntegrationTests.request.__func__)

    @classmethod
    def setUpClass(cls):
        customers.CustomersIntegrationTests.setUpClass.__func__(cls)
        # unittest skips tearDownClass if seeding fails; never leave its API alive.
        cls.addClassCleanup(lambda: cls.tearDownClass() if cls.server.poll() is None else None)
        with closing(sqlite3.connect(cls.database)) as db, db:
            # Seed 1,000 rows independently in every supported dataset.
            db.execute("PRAGMA foreign_keys=ON")
            db.execute("DELETE FROM bank_accounts")
            db.execute("DELETE FROM bank_directory")
            db.execute("DELETE FROM inventory_options")
            def insert(table, rows):
                names = list(rows[0]); marks=','.join('?' for _ in names)
                db.executemany(f"INSERT INTO {table} ({','.join(names)}) VALUES ({marks})", [tuple(r[n] for n in names) for r in rows])
            now = "2026-10-31 12:34:56.123456"
            insert("customers", [dict(id=i,customer_code=i,permanent_id=f"050809-123456-{i:04d}",
                first_name="علی " + str(i),last_name="رضایی",gender="male" if i%2 else None,
                address=None if i%2 else "",last_contact_date="2026-10-31" if i%2 else None,
                phone_numbers=json.dumps(["09120000000","۰۹۱۲"],ensure_ascii=False),
                notes="=SUM(1,2)" if i==1 else "سلام فارسی English 123\nخط دوم",created_at=now,updated_at=now,
                deleted_at=now if i==1000 else None) for i in range(1,1001)])
            insert("bank_directory", [dict(code=f"bank-{i}",name=f"بانک {i}",logo_url="/bank.svg",active=1) for i in range(1,1001)])
            insert("bank_accounts", [dict(id=i,name=f"حساب {i}",account_number=f"00123400000000000000{i}",opening_balance=987654321098765432,
                bank=f"bank-{i}",active=1,card_number=f"001234567890{i:04d}",created_at=now,updated_at=now) for i in range(1,1001)])
            insert("expense_categories", [dict(id=i,title=f"هزینه {i}",active=1,display_code=i) for i in range(1,1001)])
            insert("products", [dict(id=i,name=f"کالا {i}",product_code=f"P-{i}",unit="عدد",price=123456,stock=None if i%2 else 10,active=1,revision=1) for i in range(1,1001)])
            insert("documents", [dict(id=i,number=str(i),description="سند فارسی",date="2026-10-31",document_type="دریافت",created_at=now,revision=1) for i in range(1,1001)])
            insert("document_accounts", [dict(id=i,code=f"D-{i}",title=f"حساب سند {i}",kind="income",active=1) for i in range(1,1001)])
            insert("financial_transactions", [dict(id=i,transaction_number=i,request_id=f"finance-{i}",kind="income",customer_id=i,bank_id=i,
                title=f"دریافت {i}",account_code=f"C-{i}",amount=1000+i,transaction_date="2026-10-31" if i%2 else "2026-10-30",
                transaction_time="12:34:56",time_mode="manual",created_at=now) for i in range(1,1001)])
            insert("orders", [dict(id=i,request_id=f"order-{i}",request_fingerprint=f"fingerprint-{i}",invoice_code=f"I-{i}",order_number=f"O-{i}",
                customer_id=i,invoice_date="2026-10-31",payment_date="2026-11-01",description=None,subtotal=200000,discount=0,
                status="active" if i%2 else "completed",created_at=now,invoice_snapshot=json.dumps({"مشتری":"علی", "null":None,"list":[1,False,""]},ensure_ascii=False)) for i in range(1,1001)])
            insert("order_items", [dict(id=i,order_id=i,product_id=i,position=0,product_name=f"کالا {i}",unit="عدد",quantity=2,unit_price=100000) for i in range(1,1001)])
            insert("order_payments", [dict(id=i,order_id=i,transaction_id=i,method="cash") for i in range(1,1001)])
            insert("inventory_options", [dict(id=i,kind="categories" if i%2 else "units",name=f"گزینه {i}",active=1) for i in range(1,1001)])
            insert("members", [dict(id=i,first_name="عضو",last_name=str(i),phone="09120000000",national_id=f"001234{i:04d}",created_at=now,updated_at=now) for i in range(1,1001)])
            insert("attendance", [dict(id=i,member_id=i,date="2026-10-31",status="present" if i%2 else "absent",recorded_time="12:34:56.123456",created_at=now,updated_at=now) for i in range(1,1001)])
            insert("tasks", [dict(id=i,date="2026-10-31",title=f"کار {i}",description="شرح",status="todo" if i%2 else "completed",created_at=now) for i in range(1,1001)])

    def binary(self, name, fmt, body):
        r = Request(self.base+f"/reports/{name}/export/{fmt}", data=json.dumps(body).encode(),
            headers={"Authorization":"Bearer "+self.token,"Content-Type":"application/json"},method="POST")
        with urlopen(r,timeout=180) as response:
            self.assertEqual(response.headers["Cache-Control"],"no-store")
            return response.read()

    @staticmethod
    def xlsx_data(raw):
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=False)
        chunks=defaultdict(list)
        for kind,index,part,value in wb["داده خام"].iter_rows(values_only=True): chunks[(kind,index)].append((part,value))
        data={key:json.loads(''.join(v for _,v in sorted(parts))) for key,parts in chunks.items()}
        rows=[data[key] for key in data if key[0]=="record"]
        return wb,data[("schema",0)]["columns"],rows

    def test_01_catalog_and_schema_match_actual_database(self):
        status,items=self.request("GET","/reports")
        self.assertEqual(status,200); self.assertEqual({r["name"] for r in items},set(TABLES))
        with closing(sqlite3.connect(self.database)) as db:
            for name in TABLES:
                with self.subTest(module=name):
                    status,data=self.request("GET",f"/reports/{name}/schema")
                    self.assertEqual(status,200)
                    actual={r[1] for r in db.execute(f"PRAGMA table_info({name})")}
                    base={c["name"] for c in data["columns"] if not c["related"]}
                    self.assertEqual(base,actual)
                    self.assertTrue(all(c["label"] for c in data["columns"]))
        for name in ["users","member_faces","member_face_photos","sqlite_master","customers;DROP TABLE users"]:
            self.assertEqual(self.request("GET",f"/reports/{quote(name)}/schema")[0],404)
        self.assertEqual(self.request("GET","/reports",authenticated=False)[0],401)
        self.assertEqual(self.request("POST","/reports/customers/preview",{},authenticated=False)[0],401)
        self.assertEqual(self.request("POST","/reports/customers/export/pdf",{},authenticated=False)[0],401)

    def test_02_every_module_1000_records_in_both_formats(self):
        artifact_dir=Path(self.temp.name)/"artifacts"; artifact_dir.mkdir(exist_ok=True)
        for name in TABLES:
            with self.subTest(module=name):
                body={"include_deleted":True}
                status,data=self.request("POST",f"/reports/{name}/preview",body)
                self.assertEqual(status,200); self.assertEqual(data["total"],1000); self.assertEqual(len(data["items"]),50)
                start=time.perf_counter(); xlsx=self.binary(name,"xlsx",body); xlsx_time=time.perf_counter()-start
                wb,cols,xrows=self.xlsx_data(xlsx)
                self.assertEqual(len(xrows),1000); self.assertEqual(list(xrows[0]),data["columns"])
                sheet=wb["گزارش"]; self.assertEqual(sum(1 for _ in sheet.rows),1001)
                wb.close()
                start=time.perf_counter(); pdf=self.binary(name,"pdf",body); pdf_time=time.perf_counter()-start
                reader=PdfReader(BytesIO(pdf)); prows=json.loads(reader.attachments["report-data.json"][0])
                self.assertEqual(prows["columns"],cols); self.assertEqual(prows["records"],xrows)
                self.assertGreater(len(reader.pages),0); self.assertEqual(reader.trailer["/Root"]["/Lang"],"fa-IR")
                print(f"\nREPORT {name}: 1000 rows, XLSX {xlsx_time:.2f}s/{len(xlsx)} bytes; PDF {pdf_time:.2f}s/{len(pdf)} bytes/{len(reader.pages)} pages",flush=True)
                # Retain representative files outside the fixture for visual QA only.
                if name in {"customers","financial_transactions","orders","attendance","products"}:
                    Path(f"/tmp/karbon-reports-{name}.pdf").write_bytes(pdf)
                    Path(f"/tmp/karbon-reports-{name}.xlsx").write_bytes(xlsx)

    def test_03_selection_filters_related_fields_and_pagination(self):
        for name in TABLES:
            pk="code" if name=="bank_directory" else "id"
            one="bank-1" if pk=="code" else 1
            two="bank-2" if pk=="code" else 2
            three="bank-3" if pk=="code" else 3
            body={"columns":[pk],"ids":[one,two,three],"excluded_ids":[two],"include_deleted":True}
            with self.subTest(module=name):
                status,data=self.request("POST",f"/reports/{name}/preview",body)
                self.assertEqual(status,200); self.assertEqual(data["total"],2)
                expected=[r["values"] for r in data["items"]]
                for fmt in ["pdf","xlsx"]:
                    raw=self.binary(name,fmt,body)
                    actual=json.loads(PdfReader(BytesIO(raw)).attachments["report-data.json"][0])["records"] if fmt=="pdf" else self.xlsx_data(raw)[2]
                    self.assertEqual(actual,expected)
                self.assertEqual(self.request("POST",f"/reports/{name}/preview",{"ids":[]})[1]["total"],0)
                status,paged=self.request("POST",f"/reports/{name}/preview",{"limit":200,"offset":200,"include_deleted":True})
                self.assertEqual(status,200); self.assertEqual(len(paged["items"]),200)
        body={"columns":["id","customer_id.first_name","customer_id.permanent_id"],"filters":[{"field":"status","value":"active"},{"field":"invoice_date","op":"gte","value":"1405/08/09"},{"field":"invoice_date","op":"lte","value":"1405/08/09"}]}
        status,data=self.request("POST","/reports/orders/preview",body)
        self.assertEqual(status,200); self.assertEqual(data["total"],500)
        self.assertEqual(data["items"][0]["values"]["customer_id.first_name"],"علی 1")
        self.assertEqual(data["items"][0]["values"]["customer_id.permanent_id"],"050809-123456-0001")
        self.assertEqual(self.request("POST","/reports/customers/preview",{})[1]["total"],999)
        self.assertEqual(self.request("POST","/reports/customers/preview",{"filters":[{"field":"address","op":"is_null"}]})[1]["total"],500)
        self.assertEqual(self.request("POST","/reports/customers/preview",{"filters":[{"field":"address","value":""}]})[1]["total"],499)
        self.assertEqual(self.request("POST","/reports/customers/preview",{"search":"%' OR 1=1 --"})[1]["total"],0)
        status,options=self.request("GET","/reports/orders/options?field=status")
        self.assertEqual(status,200); self.assertEqual(set(options["values"]),{"active","completed"})
        status,options=self.request("GET","/reports/customers/options?field=first_name")
        self.assertEqual(status,200);self.assertEqual(len(options["values"]),100);self.assertTrue(options["has_more"])
        status,options=self.request("GET","/reports/orders/options?field=customer_id&search="+quote("علی 123"))
        self.assertEqual(status,200);self.assertEqual(options["choices"][0]["value"],123)
        self.assertIn("علی 123",options["choices"][0]["label"])

    def test_04_types_no_formulas_and_no_silent_omissions(self):
        raw=self.binary("customers","xlsx",{"ids":[1,2],"columns":["id","address","phone_numbers","notes","permanent_id","last_contact_date"]})
        wb,cols,rows=self.xlsx_data(raw)
        self.assertIsNone(rows[0]["address"]);self.assertEqual(rows[1]["address"],"")
        self.assertIsInstance(rows[0]["phone_numbers"],list)
        cells=list(wb["گزارش"].rows)
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        with ZipFile(BytesIO(raw)) as package:
            xml=ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
            ns={"x":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            self.assertEqual(xml.find("x:sheetViews/x:sheetView",ns).get("rightToLeft"),"1")
            self.assertEqual(xml.find("x:sheetViews/x:sheetView/x:pane",ns).get("state"),"frozen")
        self.assertEqual(cells[1][3].alignment.readingOrder,2)
        self.assertEqual(cells[1][3].font.name,"Vazirmatn")
        self.assertEqual(cells[1][3].value,"=SUM(1,2)");self.assertEqual(cells[1][3].data_type,"s")
        self.assertEqual(cells[1][4].value,"050809-123456-0001");self.assertEqual(cells[1][4].data_type,"s")
        wb.close()
        raw=self.binary("bank_accounts","xlsx",{"ids":[1],"columns":["opening_balance","active","card_number"]})
        wb,_,rows=self.xlsx_data(raw);cells=list(wb["گزارش"].rows)
        self.assertEqual(rows[0]["opening_balance"],987654321098765432)
        self.assertEqual(cells[1][0].value,"987654321098765432");self.assertTrue(cells[1][1].value);self.assertTrue(cells[1][2].value.startswith("00"));wb.close()
        for fmt in ["xlsx","pdf"]:
            raw=self.binary("customers",fmt,{"ids":[],"columns":["first_name"]})
            actual=json.loads(PdfReader(BytesIO(raw)).attachments["report-data.json"][0])["records"] if fmt=="pdf" else self.xlsx_data(raw)[2]
            self.assertEqual(actual,[])

    def test_05_invalid_requests_are_rejected(self):
        cases=[{"columns":["nonexistent"]},{"columns":[]},{"columns":["id","id"]},
               {"filters":[{"field":"password_hash","value":"x"}]},
               {"filters":[{"field":"created_at","op":"gte","value":"not-date"}]},
               {"filters":[{"field":"id","value":"1.2"}]},
               {"filters":[{"field":"first_name","op":"gte","value":"علی"}]},
               {"filters":[{"field":"created_at","op":"gte","value":"2026-11-01"},{"field":"created_at","op":"lte","value":"2026-10-01"}]},
               {"limit":201},{"offset":-1},{"tables":["users"]}]
        for body in cases:
            with self.subTest(body=body): self.assertEqual(self.request("POST","/reports/customers/preview",body)[0],422)
        self.assertEqual(self.request("POST","/reports/customers/export/csv",{})[0],422)

    def test_06_long_multiline_values_and_schema_unchanged(self):
        with closing(sqlite3.connect(self.database)) as db:
            before=list(db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name"))
            original=db.execute("SELECT notes FROM customers WHERE id=1").fetchone()[0]
            long="متن بلند فارسی English 123\n"*1800
            db.execute("UPDATE customers SET notes=? WHERE id=1",(long,));db.commit()
        try:
            for fmt in ["pdf","xlsx"]:
                raw=self.binary("customers",fmt,{"columns":["notes"],"ids":[1]})
                rows=json.loads(PdfReader(BytesIO(raw)).attachments["report-data.json"][0])["records"] if fmt=="pdf" else self.xlsx_data(raw)[2]
                self.assertEqual(rows,[{"notes":long}])
            with closing(sqlite3.connect(self.database)) as db:
                after=list(db.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name"))
                self.assertEqual(before,after)
        finally:
            with closing(sqlite3.connect(self.database)) as db,db:db.execute("UPDATE customers SET notes=? WHERE id=1",(original,))

    def test_07_module_specific_filters_match_sqlite(self):
        with closing(sqlite3.connect(self.database)) as db:
            db.row_factory=sqlite3.Row
            for name in TABLES:
                actual={r[1]:r[2] for r in db.execute(f"PRAGMA table_info({name})")}
                first=db.execute(f"SELECT * FROM {name} LIMIT 1").fetchone()
                fields=[n for n in actual if n in {"status","kind","active","document_type","gender","customer_id","bank_id","member_id","order_id","product_id","transaction_id"}]
                fields += [n for n,t in actual.items() if t in {"DATE","DATETIME"}][:1]
                for field in fields:
                    with self.subTest(module=name,field=field):
                        value=first[field]
                        if value is None:continue
                        expected=db.execute(f"SELECT count(*) FROM {name} WHERE {field}=?",(value,)).fetchone()[0]
                        if actual[field]=="BOOLEAN":value=bool(value)
                        request={"include_deleted":True,"filters":[{"field":field,"value":value}]}
                        status,result=self.request("POST",f"/reports/{name}/preview",request)
                        self.assertEqual(status,200,result);self.assertEqual(result["total"],expected)

    def test_08_export_releases_read_lock_and_retains_snapshot(self):
        with closing(sqlite3.connect(self.database)) as db:
            original=db.execute("SELECT notes FROM customers WHERE id=2").fetchone()[0]
        try:
            with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
                future=pool.submit(self.binary,"customers","pdf",{"columns":["id","first_name","notes"],"include_deleted":True})
                time.sleep(.5)
                # Rendering 1,000 Persian rows must not keep SQLite locked.
                with closing(sqlite3.connect(self.database,timeout=.25)) as db,db:
                    db.execute("UPDATE customers SET notes=? WHERE id=2",("ویرایش همزمان",))
                raw=future.result(timeout=180)
                rows=json.loads(PdfReader(BytesIO(raw)).attachments["report-data.json"][0])["records"]
                self.assertEqual(next(r for r in rows if r["id"]==2)["notes"],original)
        finally:
            with closing(sqlite3.connect(self.database)) as db,db:
                db.execute("UPDATE customers SET notes=? WHERE id=2",(original,))

    def test_09_all_related_fields_and_large_selections(self):
        for name in TABLES:
            with self.subTest(module=name):
                status,schema=self.request("GET",f"/reports/{name}/schema")
                columns=[c["name"] for c in schema["columns"]]
                key="bank-1" if name=="bank_directory" else 1
                body={"columns":columns,"ids":[key]}
                status,preview=self.request("POST",f"/reports/{name}/preview",body)
                self.assertEqual(status,200)
                xlsx=self.binary(name,"xlsx",body)
                wb,_,xrows=self.xlsx_data(xlsx);wb.close()
                pdf=self.binary(name,"pdf",body)
                prows=json.loads(PdfReader(BytesIO(pdf)).attachments["report-data.json"][0])["records"]
                self.assertEqual(prows,xrows)
                self.assertEqual(list(xrows[0]),columns)
        ids=list(range(1,1001))
        status,data=self.request("POST","/reports/customers/preview",{"include_deleted":True,"ids":ids,"columns":["id"]})
        self.assertEqual(status,200);self.assertEqual(data["total"],1000)
        status,data=self.request("POST","/reports/customers/preview",{"include_deleted":True,"excluded_ids":ids,"columns":["id"]})
        self.assertEqual(status,200);self.assertEqual(data["total"],0)


if __name__ == "__main__": unittest.main()
